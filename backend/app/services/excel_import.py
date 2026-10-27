from datetime import date, datetime
from typing import BinaryIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.license_item import LicenseItem
from app.schemas.license_item import UploadResult

# Maps the exact column header text in the uploaded workbook to the
# LicenseItem attribute it should be stored in. Columns not listed here
# (e.g. "Folder Name", "License Type New", "Quantity") are ignored.
COLUMN_MAP: dict[str, str] = {
    "Serial Number": "serial_number",
    "Parent Customer": "parent_customer",
    "Customer Name": "customer_name",
    "License Key": "license_key",
    "License Type": "license_type",
    "Product Name": "product_name",
    "Quantity - ENT": "quantity_ent",
    "Contract Status": "contract_status",
    "Start Date": "start_date",
    "End Date": "end_date",
    "ERP Site ID": "erp_site_id",
    "Enduser Account Name": "enduser_account_name",
    "Part Number": "part_number",
    "Division": "division",
}

REQUIRED_HEADERS = set(COLUMN_MAP.keys())

DATE_FORMATS = (
    "%m/%d/%y %I:%M:%S %p",
    "%m/%d/%Y %I:%M:%S %p",
    "%m/%d/%y",
    "%m/%d/%Y",
    "%Y-%m-%d",
)


def _parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_quantity(value) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(float(str(value).replace(",", "").strip()))
    except ValueError:
        return 0


def _clean_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_excel(db: Session, file_obj: BinaryIO) -> UploadResult:
    workbook = load_workbook(file_obj, data_only=True, read_only=True)
    sheet = workbook.active

    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        return UploadResult(inserted=0, updated=0, skipped=0, errors=["빈 파일입니다."])

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    col_index = {name: idx for idx, name in enumerate(headers)}

    missing = REQUIRED_HEADERS - set(col_index.keys())
    if missing:
        raise ValueError(f"필수 컬럼이 없습니다: {', '.join(sorted(missing))}")

    inserted = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    existing_by_serial = {item.serial_number: item for item in db.query(LicenseItem).all()}

    def cell(row: tuple, header: str):
        idx = col_index[header]
        return row[idx] if idx < len(row) else None

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(c is None for c in row):
            continue

        serial_number = _clean_str(cell(row, "Serial Number"))
        if not serial_number:
            skipped += 1
            errors.append(f"{row_number}행: Serial Number가 없어 건너뜀")
            continue

        values = {
            "serial_number": serial_number,
            "parent_customer": _clean_str(cell(row, "Parent Customer")),
            "customer_name": _clean_str(cell(row, "Customer Name")),
            "license_key": _clean_str(cell(row, "License Key")),
            "license_type": _clean_str(cell(row, "License Type")),
            "product_name": _clean_str(cell(row, "Product Name")),
            "quantity_ent": _parse_quantity(cell(row, "Quantity - ENT")),
            "contract_status": _clean_str(cell(row, "Contract Status")),
            "start_date": _parse_date(cell(row, "Start Date")),
            "end_date": _parse_date(cell(row, "End Date")),
            "erp_site_id": _clean_str(cell(row, "ERP Site ID")),
            "enduser_account_name": _clean_str(cell(row, "Enduser Account Name")),
            "part_number": _clean_str(cell(row, "Part Number")),
            "division": _clean_str(cell(row, "Division")),
        }

        existing = existing_by_serial.get(serial_number)
        if existing is not None:
            for key, value in values.items():
                if key == "serial_number":
                    continue
                setattr(existing, key, value)
            updated += 1
        else:
            item = LicenseItem(**values)
            db.add(item)
            existing_by_serial[serial_number] = item
            inserted += 1

    db.commit()
    return UploadResult(inserted=inserted, updated=updated, skipped=skipped, errors=errors)
